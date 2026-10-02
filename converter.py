"""Convert documents of many formats to PDF.

Strategy: if LibreOffice (soffice) is installed it is used for office formats
(best fidelity: doc, docx, odt, rtf, xls, xlsx, ppt, pptx, ...). Otherwise
pure-Python fallbacks handle docx, xlsx, csv, txt/md/code, html, and images.
"""
import csv, html, os, re, shutil, subprocess, sys, tempfile
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, Preformatted, PageBreak)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tiff", ".tif", ".webp"}
TEXT_EXT = {".txt", ".md", ".log", ".json", ".xml", ".py", ".js", ".java",
            ".c", ".cpp", ".css", ".yaml", ".yml", ".ini", ".sql", ".sh"}
OFFICE_EXT = {".doc", ".docx", ".odt", ".rtf", ".xls", ".xlsx", ".ods",
              ".ppt", ".pptx", ".odp", ".wpd", ".pages"}
SUPPORTED = IMAGE_EXT | TEXT_EXT | OFFICE_EXT | {".csv", ".html", ".htm", ".pdf"}

styles = getSampleStyleSheet()


def _soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def _libreoffice(src: Path, out: Path):
    with tempfile.TemporaryDirectory() as d:
        subprocess.run([_soffice(), "--headless", "--convert-to", "pdf",
                        "--outdir", d, str(src)], check=True,
                       capture_output=True, timeout=180)
        shutil.move(os.path.join(d, src.stem + ".pdf"), out)


def _build(out, story, pagesize=A4):
    SimpleDocTemplate(str(out), pagesize=pagesize, leftMargin=2*cm,
                      rightMargin=2*cm, topMargin=2*cm,
                      bottomMargin=2*cm).build(story or [Spacer(1, 1)])


def _table(rows, width):
    if not rows:
        return Spacer(1, 1)
    n = max(len(r) for r in rows)
    cell = styles["BodyText"].clone("cell", fontSize=8, leading=10)
    data = [[Paragraph(escape(str(c if c is not None else "")), cell)
             for c in list(r) + [""] * (n - len(r))] for r in rows]
    t = Table(data, colWidths=[width / n] * n, repeatRows=1)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _docx(src, out):
    import docx
    d = docx.Document(str(src))
    story, width = [], A4[0] - 4*cm
    body = d.element.body
    for el in body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(el, d)
            sname = (p.style.name or "").lower()
            parts = []
            for r in p.runs:
                t = escape(r.text)
                if r.bold: t = f"<b>{t}</b>"
                if r.italic: t = f"<i>{t}</i>"
                if r.underline: t = f"<u>{t}</u>"
                parts.append(t)
            text = "".join(parts)
            if not text.strip():
                story.append(Spacer(1, 6)); continue
            if sname.startswith("title"): st = styles["Title"]
            elif sname.startswith("heading"):
                lvl = re.findall(r"\d", sname)
                st = styles[f"Heading{min(int(lvl[0]) if lvl else 1, 6)}"]
            elif "list" in sname: text, st = "• " + text, styles["BodyText"]
            else: st = styles["BodyText"]
            story.append(Paragraph(text, st))
        elif tag == "tbl":
            t = docx.table.Table(el, d)
            story += [_table([[c.text for c in r.cells] for r in t.rows], width),
                      Spacer(1, 8)]
    _build(out, story)


def _xlsx(src, out):
    import openpyxl
    wb = openpyxl.load_workbook(str(src), data_only=True, read_only=True)
    ps = landscape(A4); story = []
    for i, ws in enumerate(wb.worksheets):
        if i: story.append(PageBreak())
        rows = [r for r in ws.iter_rows(values_only=True)
                if any(c is not None for c in r)]
        story += [Paragraph(escape(ws.title), styles["Heading2"]),
                  _table(rows, ps[0] - 4*cm)]
    _build(out, story, ps)


def _csv(src, out):
    with open(src, newline="", encoding="utf-8", errors="replace") as f:
        rows = list(csv.reader(f))
    ps = landscape(A4)
    _build(out, [_table(rows, ps[0] - 4*cm)], ps)


def _text(src, out):
    txt = Path(src).read_text(encoding="utf-8", errors="replace")
    mono = styles["Code"].clone("mono", fontSize=8, leading=10)
    # wrap long lines
    lines = []
    for ln in txt.expandtabs(4).splitlines():
        lines += [ln[i:i+100] for i in range(0, max(len(ln), 1), 100)]
    _build(out, [Preformatted("\n".join(lines), mono)])


def _html(src, out):
    raw = Path(src).read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"(?is)<(script|style|head).*?</\1>", "", raw)
    story = []
    for m in re.finditer(r"(?is)<(h[1-6]|p|li|div|td|pre)[^>]*>(.*?)</\1>", raw):
        tag, inner = m.group(1).lower(), m.group(2)
        text = escape(html.unescape(re.sub(r"<[^>]+>", "", inner))).strip()
        if not text: continue
        st = styles[f"Heading{tag[1]}"] if tag[0] == "h" else styles["BodyText"]
        story.append(Paragraph(("• " if tag == "li" else "") + text, st))
    if not story:
        text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
        story = [Paragraph(escape(t), styles["BodyText"])
                 for t in text.split("\n") if t.strip()]
    _build(out, story)


def _images(srcs, out):
    frames = []
    for s in srcs:
        im = Image.open(s)
        try:
            n = getattr(im, "n_frames", 1)
            for i in range(n if Path(s).suffix.lower() in {".tif", ".tiff"} else 1):
                im.seek(i); frames.append(im.convert("RGB"))
        except EOFError:
            pass
    frames[0].save(out, "PDF", save_all=True, append_images=frames[1:],
                   resolution=100)


def convert(src, out=None) -> Path:
    src = Path(src)
    ext = src.suffix.lower()
    out = Path(out) if out else src.with_suffix(".pdf")
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported format: {ext or '(none)'}")
    if ext == ".pdf":
        shutil.copy(src, out)
    elif ext in IMAGE_EXT:
        _images([src], out)
    elif _soffice() and ext not in TEXT_EXT:
        _libreoffice(src, out)
    elif ext == ".docx": _docx(src, out)
    elif ext == ".xlsx": _xlsx(src, out)
    elif ext == ".csv": _csv(src, out)
    elif ext in {".html", ".htm"}: _html(src, out)
    elif ext in TEXT_EXT: _text(src, out)
    else:
        raise RuntimeError(f"{ext} needs LibreOffice installed "
                           "(apt install libreoffice)")
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("usage: python converter.py FILE [FILE ...]")
    for f in sys.argv[1:]:
        try: print(f"✔ {f} -> {convert(f)}")
        except Exception as e: print(f"✘ {f}: {e}")
